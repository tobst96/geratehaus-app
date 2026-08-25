"""Tests für Dienstbuch-Planer Phase 2-4: Feiertage, Excel-Export/-Import,
Divera-Übertragung (mit gemocktem Divera-Client)."""

import io
from datetime import date, time

from openpyxl import load_workbook

from app.core.security import hash_secret
from app.models.person import Person
from app.schemas.dienstbuch_planer import PlanTerminAnlegen
from app.services import dienstbuch_planer_service, divera_planer_service, feiertag_service
from app.services.config_service import config_service
from app.services.dienstbuch_planer_excel_service import jahres_export_xlsx, jahres_import_xlsx


async def _token(client, db, name="gf"):
    db.add(Person(name=name, email=name, passwort_hash=hash_secret("geheim123"), gruppenfuehrer_rolle="gruppenfuehrer"))
    await db.commit()
    login = await client.post(
        "/api/v1/auth/gruppenfuehrer/login", data={"username": name, "password": "geheim123"}
    )
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


# --- Phase 2: Feiertage ---------------------------------------------------------


def test_ostersonntag_bekannte_jahre():
    assert feiertag_service.ostersonntag(2025) == date(2025, 4, 20)
    assert feiertag_service.ostersonntag(2026) == date(2026, 4, 5)
    assert feiertag_service.ostersonntag(2027) == date(2027, 3, 28)


def test_berechne_feiertage_bundesweit_und_landesspezifisch():
    bundesweit = feiertag_service.berechne_feiertage(2026, "")
    namen = {f.name for f in bundesweit}
    assert "Neujahr" in namen
    assert "Karfreitag" in namen
    assert "Allerheiligen" not in namen  # landesspezifisch

    bayern = feiertag_service.berechne_feiertage(2026, "BY")
    namen_by = {f.name for f in bayern}
    assert "Allerheiligen" in namen_by
    assert "Fronleichnam" in namen_by
    # Karfreitag 2026 = Ostern (5.4.) - 2 Tage
    karfreitag = next(f for f in bayern if f.name == "Karfreitag")
    assert karfreitag.datum == date(2026, 4, 3)


def test_buss_und_bettag_nur_sachsen():
    sachsen = {f.name for f in feiertag_service.berechne_feiertage(2026, "SN")}
    assert "Buß- und Bettag" in sachsen
    bbtag = next(
        f for f in feiertag_service.berechne_feiertage(2026, "SN") if f.name == "Buß- und Bettag"
    )
    assert bbtag.datum.weekday() == 2  # Mittwoch
    assert bbtag.datum < date(2026, 11, 23)


async def test_feiertage_endpunkt_mit_manuellem_eintrag(client, db):
    await config_service.set(db, "dienstbuch_planer_bundesland", "BY")
    h = await _token(client, db)

    neu = await client.post(
        "/api/v1/dienstbuch-planer/feiertage",
        json={"datum": "2026-07-11", "name": "Sommerfest (blockiert)"},
        headers=h,
    )
    assert neu.status_code == 201
    feiertag_id = neu.json()["id"]

    liste = await client.get("/api/v1/dienstbuch-planer/feiertage?jahr=2026", headers=h)
    assert liste.status_code == 200
    eintraege = liste.json()
    assert any(e["name"] == "Sommerfest (blockiert)" and e["quelle"] == "manuell" for e in eintraege)
    assert any(e["name"] == "Allerheiligen" and e["quelle"] == "regel" for e in eintraege)

    geloescht = await client.delete(f"/api/v1/dienstbuch-planer/feiertage/{feiertag_id}", headers=h)
    assert geloescht.status_code == 204


# --- Phase 3: Excel -------------------------------------------------------------


async def test_excel_export_und_reimport_roundtrip(db):
    await dienstbuch_planer_service.termin_anlegen(
        db,
        PlanTerminAnlegen(
            titel="Unterweisung UVV",
            zieldatum=date(2026, 1, 28),
            uhrzeit=time(19, 30),
            endzeit=time(21, 0),
        ),
        "Tester",
    )

    inhalt = await jahres_export_xlsx(db, 2026)
    wb = load_workbook(io.BytesIO(inhalt))
    assert len(wb.sheetnames) == 12
    ws = wb["Januar"]
    zeilen = [[c.value for c in row] for row in ws.iter_rows()]
    flach = [str(z) for zeile in zeilen for z in zeile if z]
    assert any("Unterweisung UVV" in z for z in flach)
    assert any("Gerätehaus.app" in z for z in flach)

    # Re-Import derselben Datei: bestehender Termin wird übersprungen.
    ergebnis = await jahres_import_xlsx(db, inhalt, 2026, "Tester")
    assert ergebnis.angelegt == 0
    assert ergebnis.uebersprungen == 1
    assert ergebnis.fehler == []


async def test_excel_import_legt_neue_termine_an(db):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Datum", "Beginn", "Ende", "Titel", "Beschreibung", "Kategorien", "Status"])
    ws.append(["04.03.2026", "19:00", "21:00", "Übungsdienst", "Knoten & Stiche", "", ""])
    ws.append(["kaputt", "", "", "Ohne Datum", "", "", ""])
    puffer = io.BytesIO()
    wb.save(puffer)

    ergebnis = await jahres_import_xlsx(db, puffer.getvalue(), 2026, "Tester")
    assert ergebnis.angelegt == 1
    assert len(ergebnis.fehler) == 1

    termine = await dienstbuch_planer_service.liste_termine(db, 2026)
    neu = next(t for t in termine if t.titel == "Übungsdienst")
    assert neu.zieldatum == date(2026, 3, 4)
    assert neu.uhrzeit == time(19, 0)
    assert neu.endzeit == time(21, 0)
    assert neu.status == "entwurf"


# --- Phase 4: Divera ------------------------------------------------------------


async def test_divera_uebertragung_baut_korrektes_event(db, monkeypatch):
    await config_service.set(db, "divera_api_key", "test-key")
    termin = await dienstbuch_planer_service.termin_anlegen(
        db,
        PlanTerminAnlegen(
            titel="Ortskommandositzung",
            beschreibung="Themen laut Einladung",
            zieldatum=date(2026, 9, 10),
            uhrzeit=time(19, 30),
            endzeit=time(21, 0),
        ),
        "Tester",
    )

    aufrufe: list[dict] = []

    async def fake_erstelle_termin(api_key, event, reminder=None):
        aufrufe.append({"api_key": api_key, "event": event, "reminder": reminder})
        return True, ""

    from app.services import divera_client

    monkeypatch.setattr(divera_client, "erstelle_termin", fake_erstelle_termin)

    ergebnisse = await divera_planer_service.uebertrage_termine(
        db, [termin.id], [3, 7], 60, True, "Chefin"
    )
    assert len(ergebnisse) == 1 and ergebnisse[0].ok

    aufruf = aufrufe[0]
    assert aufruf["api_key"] == "test-key"
    event = aufruf["event"]
    assert event["title"] == "Ortskommandositzung"
    assert event["notification_type"] == 3
    assert event["group"] == [3, 7]
    assert event["fullday"] is False
    assert event["ts_end"] - event["ts_start"] == 90 * 60
    assert aufruf["reminder"]["ts"] == event["ts_start"] - 3600

    ereignisse = await dienstbuch_planer_service.termin_ereignisse(db, termin.id)
    assert any(e.typ == "divera_uebertragen" and e.akteur_name == "Chefin" for e in ereignisse)


async def test_divera_uebertragung_ohne_key_liefert_fehler(db):
    await config_service.set(db, "divera_api_key", "")
    ergebnisse = await divera_planer_service.uebertrage_termine(db, [1], [], None, True, None)
    assert len(ergebnisse) == 1
    assert ergebnisse[0].ok is False
    assert "API-Key" in ergebnisse[0].fehler


async def test_divera_uebertragung_platzhalter_abgelehnt(db, monkeypatch):
    await config_service.set(db, "divera_api_key", "test-key")
    from app.schemas.dienstbuch_planer import PlanPlatzhalterAnlegen

    platzhalter = await dienstbuch_planer_service.platzhalter_anlegen(
        db, PlanPlatzhalterAnlegen(titel="Noch offen", jahr=2026), None
    )
    ergebnisse = await divera_planer_service.uebertrage_termine(
        db, [platzhalter.id], [], None, True, None
    )
    assert ergebnisse[0].ok is False
    assert "Platzhalter" in ergebnisse[0].fehler


async def test_geseedete_feiertage_sind_loeschbar_und_kommen_nicht_wieder(client, db):
    """Nutzerwunsch 25.08.2026: auch gesetzliche Feiertage müssen löschbar
    sein - Seed passiert einmalig, gelöschte Tage tauchen bei erneutem Abruf
    NICHT wieder auf (Seed-Marker in der Config, keine Existenz-Heuristik)."""
    await config_service.set(db, "dienstbuch_planer_bundesland", "BY")
    h = await _token(client, db)

    liste = (await client.get("/api/v1/dienstbuch-planer/feiertage?jahr=2028", headers=h)).json()
    neujahr = next(e for e in liste if e["name"] == "Neujahr")
    assert neujahr["quelle"] == "regel"
    assert neujahr["id"] is not None  # geseedet -> DB-Zeile -> löschbar

    geloescht = await client.delete(
        f"/api/v1/dienstbuch-planer/feiertage/{neujahr['id']}", headers=h
    )
    assert geloescht.status_code == 204

    erneut = (await client.get("/api/v1/dienstbuch-planer/feiertage?jahr=2028", headers=h)).json()
    assert not any(e["name"] == "Neujahr" for e in erneut)


async def test_divera_info_nur_mit_modul_und_key(client, db, monkeypatch):
    h = await _token(client, db)

    async def fake_hole_gruppen(api_key):
        return [{"id": 4, "name": "Aktive"}, {"id": 9, "name": "Jugend"}]

    from app.services import divera_client

    monkeypatch.setattr(divera_client, "hole_gruppen", fake_hole_gruppen)

    # Modul aus -> inaktiv, keine Gruppen.
    await config_service.set(db, "modul_divera_aktiv", False)
    await config_service.set(db, "divera_api_key", "test-key")
    info = (await client.get("/api/v1/dienstbuch-planer/divera-info", headers=h)).json()
    assert info["aktiv"] is False and info["gruppen"] == []

    # Modul an, aber kein Key -> inaktiv.
    await config_service.set(db, "modul_divera_aktiv", True)
    await config_service.set(db, "divera_api_key", "")
    info = (await client.get("/api/v1/dienstbuch-planer/divera-info", headers=h)).json()
    assert info["aktiv"] is False

    # Modul an + Key -> aktiv mit Gruppen aus der API.
    await config_service.set(db, "divera_api_key", "test-key")
    info = (await client.get("/api/v1/dienstbuch-planer/divera-info", headers=h)).json()
    assert info["aktiv"] is True
    assert info["gruppen"] == [{"id": 4, "name": "Aktive"}, {"id": 9, "name": "Jugend"}]


async def test_vorlage_loeschen_endgueltig_termine_bleiben(client, db):
    h = await _token(client, db)
    vorlage = await client.post(
        "/api/v1/dienstbuch-planer/vorlagen",
        json={
            "titel": "Löschtest",
            "wiederholungstyp": "jaehrlich",
            "wochentag": 2,
            "kalenderwoche": 10,
            "startdatum": "2020-01-01",
        },
        headers=h,
    )
    vorlage_id = vorlage.json()["id"]
    await client.post("/api/v1/dienstbuch-planer/termine/jahr/2026/sicherstellen", headers=h)

    geloescht = await client.delete(f"/api/v1/dienstbuch-planer/vorlagen/{vorlage_id}", headers=h)
    assert geloescht.status_code == 204

    vorlagen = (await client.get("/api/v1/dienstbuch-planer/vorlagen", headers=h)).json()
    assert not any(v["id"] == vorlage_id for v in vorlagen)

    # Erzeugte Termine bleiben als Einzeltermine erhalten (FK SET NULL).
    termine = (await client.get("/api/v1/dienstbuch-planer/termine?jahr=2026", headers=h)).json()
    uebrig = [t for t in termine if t["titel"] == "Löschtest"]
    assert len(uebrig) == 1
    assert uebrig[0]["vorlage_id"] is None
