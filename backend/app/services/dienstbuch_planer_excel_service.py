"""Excel-Export/-Import fürs Dienstbuch-Planer-Jahr (Phase 3).

Export: ein Sheet pro Monat, mit Gerätehaus.app-Branding im Kopf und einem
QR-Code je Monat, der den Planer direkt im passenden Monat öffnet
(`/gruppenfuehrer/dienstbuch-planer?jahr=&monat=` - Login/Berechtigung
verlangt die Seite selbst).

Import: liest dieselbe Tabellenstruktur zeilenweise ein (Fehler werden pro
Zeile gesammelt statt abzubrechen - Muster wie beim Personen-CSV-Import) und
legt Entwurfs-Termine an; Zeilen, deren (Datum, Titel) im Jahr bereits
existiert, werden übersprungen (idempotenter Re-Import).
"""

import io
from dataclasses import dataclass
from datetime import date, datetime, time

import segno
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.dienstbuch_planer import DienstbuchPlanTermin
from app.services import dienstbuch_planer_service, feiertag_service
from app.services.config_service import config_service

MONATSNAMEN = [
    "Januar", "Februar", "März", "April", "Mai", "Juni",
    "Juli", "August", "September", "Oktober", "November", "Dezember",
]

_SPALTEN = ["Datum", "Beginn", "Ende", "Titel", "Beschreibung", "Kategorien", "Status"]
_WERBUNG = "Erstellt mit Gerätehaus.app – der selbst gehosteten Gerätehaus-Verwaltung"


def _zeit_str(t: time | None) -> str:
    return t.strftime("%H:%M") if t else ""


async def jahres_export_xlsx(db: AsyncSession, jahr: int) -> bytes:
    """Jahres-Export: pro Monat ein Kalenderraster-Blatt (Mo-So, Termine und
    Feiertage in den Tageszellen - wie die Kalenderansicht der App) plus ein
    flaches "Terminliste"-Blatt, das der Re-Import liest."""
    import calendar as _calendar

    from openpyxl.styles import Border, PatternFill, Side

    organisation = str(await config_service.get(db, "organisation_name", "Meine Feuerwehr"))
    basis_url = str(await config_service.get(db, "oeffentliche_basis_url", "")).rstrip("/")
    termine = await dienstbuch_planer_service.liste_termine(db, jahr)
    feiertage = await feiertag_service.feiertage_fuer_jahr(db, jahr)
    feiertag_je_tag = {f.datum: f.name for f in feiertage}
    termine_je_tag: dict[date, list] = {}
    for termin in termine:
        if termin.zieldatum is not None:
            termine_je_tag.setdefault(termin.zieldatum, []).append(termin)

    duenn = Side(style="thin", color="CCCCCC")
    rand = Border(left=duenn, right=duenn, top=duenn, bottom=duenn)
    kopf_fill = PatternFill("solid", fgColor="F3F4F6")
    feiertag_fill = PatternFill("solid", fgColor="FDE8E8")
    wochentage = ["Mo.", "Di.", "Mi.", "Do.", "Fr.", "Sa.", "So."]

    wb = Workbook()
    wb.remove(wb.active)

    for monat in range(1, 13):
        ws = wb.create_sheet(MONATSNAMEN[monat - 1])
        ws.merge_cells("A1:G1")
        ws["A1"] = f"{organisation} – Jahresdienstplan {jahr} – {MONATSNAMEN[monat - 1]}"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A2:G2")
        ws["A2"] = _WERBUNG
        ws["A2"].font = Font(italic=True, size=9)

        # QR-Code rechts NEBEN dem Kalenderraster (Spalte I) - Beschriftung in
        # I1, Bild ab I2 verankert, damit nichts überdeckt wird.
        if basis_url:
            link = f"{basis_url}/gruppenfuehrer/dienstbuch-planer?jahr={jahr}&monat={monat}"
            png = io.BytesIO()
            segno.make(link, error="m").save(png, kind="png", scale=3, border=2)
            png.seek(0)
            ws["I1"] = "Monat in der App öffnen:"
            ws["I1"].font = Font(size=9)
            ws.add_image(XlsxImage(png), "I2")

        # Wochentags-Kopf (Zeile 4) + Wochenzeilen des Monatsrasters.
        for spalte, name in enumerate(wochentage, start=1):
            zelle = ws.cell(row=4, column=spalte, value=name)
            zelle.font = Font(bold=True)
            zelle.fill = kopf_fill
            zelle.border = rand
            zelle.alignment = Alignment(horizontal="center")

        raster = _calendar.Calendar(firstweekday=0).monthdatescalendar(jahr, monat)
        zeile = 5
        for woche in raster:
            for spalte, tag in enumerate(woche, start=1):
                zelle = ws.cell(row=zeile, column=spalte)
                zelle.border = rand
                zelle.alignment = Alignment(wrap_text=True, vertical="top")
                if tag.month != monat:
                    zelle.value = str(tag.day)
                    zelle.font = Font(size=9, color="BBBBBB")
                    continue
                teile = [str(tag.day)]
                if tag in feiertag_je_tag:
                    zelle.fill = feiertag_fill
                    teile.append(feiertag_je_tag[tag])
                for termin in termine_je_tag.get(tag, []):
                    zeit = f"{_zeit_str(termin.uhrzeit)}–{_zeit_str(termin.endzeit)} " if termin.endzeit else (
                        f"{_zeit_str(termin.uhrzeit)} " if termin.uhrzeit else ""
                    )
                    stern = "* " if termin.status != "bestaetigt" else ""
                    teile.append(f"{stern}{zeit}{termin.titel}")
                zelle.value = "\n".join(teile)
            ws.row_dimensions[zeile].height = 58
            zeile += 1

        legende = ws.cell(row=zeile + 1, column=1, value="* = Entwurf · rot hinterlegt = Feiertag")
        legende.font = Font(size=9, italic=True)

        for spalte in range(1, 8):
            ws.column_dimensions[get_column_letter(spalte)].width = 20
        ws["A1"].alignment = Alignment(horizontal="left")

    # Flaches Listen-Blatt: Quelle für den Re-Import (Kopfzeile "Datum ... Titel").
    ws = wb.create_sheet("Terminliste")
    ws.merge_cells("A1:G1")
    ws["A1"] = f"{organisation} – Terminliste {jahr} (für den Re-Import in Gerätehaus.app)"
    ws["A1"].font = Font(bold=True, size=12)
    for spalte, titel in enumerate(_SPALTEN, start=1):
        zelle = ws.cell(row=3, column=spalte, value=titel)
        zelle.font = Font(bold=True)
    zeile = 4
    for termin in termine:
        if termin.zieldatum is None:
            continue
        ws.cell(row=zeile, column=1, value=termin.zieldatum.strftime("%d.%m.%Y"))
        ws.cell(row=zeile, column=2, value=_zeit_str(termin.uhrzeit))
        ws.cell(row=zeile, column=3, value=_zeit_str(termin.endzeit))
        ws.cell(row=zeile, column=4, value=termin.titel)
        ws.cell(row=zeile, column=5, value=termin.beschreibung or "")
        ws.cell(row=zeile, column=6, value=", ".join(k.name for k in termin.kategorien))
        ws.cell(row=zeile, column=7, value="Bestätigt" if termin.status == "bestaetigt" else "Entwurf")
        zeile += 1
    for spalte, breite in enumerate((12, 8, 8, 32, 40, 24, 12), start=1):
        ws.column_dimensions[get_column_letter(spalte)].width = breite

    puffer = io.BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


@dataclass
class ImportErgebnis:
    angelegt: int
    uebersprungen: int
    fehler: list[dict]


def _datum_parsen(wert: object) -> date | None:
    if isinstance(wert, datetime):
        return wert.date()
    if isinstance(wert, date):
        return wert
    if isinstance(wert, str) and wert.strip():
        for muster in ("%d.%m.%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(wert.strip(), muster).date()
            except ValueError:
                continue
    return None


def _zeit_parsen(wert: object) -> time | None:
    if isinstance(wert, time):
        return wert
    if isinstance(wert, datetime):
        return wert.time()
    if isinstance(wert, str) and wert.strip():
        for muster in ("%H:%M", "%H:%M:%S"):
            try:
                return datetime.strptime(wert.strip(), muster).time()
            except ValueError:
                continue
    return None


async def jahres_import_xlsx(
    db: AsyncSession, inhalt: bytes, jahr: int, akteur_name: str | None
) -> ImportErgebnis:
    """Liest alle Sheets der Export-Struktur ein (Kopfzeile "Datum | Beginn |
    Ende | Titel | ..."), legt fehlende Termine als Entwurf an."""
    try:
        wb = load_workbook(io.BytesIO(inhalt), read_only=True, data_only=True)
    except Exception:
        return ImportErgebnis(angelegt=0, uebersprungen=0, fehler=[{"zeile": 0, "fehler": "Keine gültige Excel-Datei (.xlsx)."}])

    bestehende = {
        (t.zieldatum, t.titel)
        for t in (
            await db.execute(
                select(DienstbuchPlanTermin).where(
                    DienstbuchPlanTermin.jahr == jahr, DienstbuchPlanTermin.zieldatum.isnot(None)
                )
            )
        ).scalars().all()
    }

    angelegt = 0
    uebersprungen = 0
    fehler: list[dict] = []

    from app.schemas.dienstbuch_planer import PlanTerminAnlegen

    for ws in wb.worksheets:
        kopf_gefunden = False
        for index, zeile in enumerate(ws.iter_rows(values_only=True), start=1):
            werte = list(zeile) + [None] * (7 - len(zeile))
            if not kopf_gefunden:
                if werte[0] == "Datum" and werte[3] == "Titel":
                    kopf_gefunden = True
                continue
            if all(w in (None, "") for w in werte[:5]):
                continue
            if isinstance(werte[0], str) and werte[0].strip().startswith("Feiertage"):
                break

            datum = _datum_parsen(werte[0])
            titel = str(werte[3] or "").strip()
            if datum is None or not titel:
                fehler.append(
                    {"zeile": index, "fehler": f"{ws.title}: Datum oder Titel fehlt/ungültig."}
                )
                continue
            if datum.year != jahr:
                fehler.append(
                    {"zeile": index, "fehler": f"{ws.title}: Datum {datum} liegt nicht im Jahr {jahr}."}
                )
                continue
            if (datum, titel) in bestehende:
                uebersprungen += 1
                continue

            beginn = _zeit_parsen(werte[1])
            ende = _zeit_parsen(werte[2]) if beginn else None
            await dienstbuch_planer_service.termin_anlegen(
                db,
                PlanTerminAnlegen(
                    titel=titel,
                    beschreibung=str(werte[4]).strip() if werte[4] else None,
                    zieldatum=datum,
                    uhrzeit=beginn,
                    endzeit=ende,
                ),
                akteur_name,
            )
            bestehende.add((datum, titel))
            angelegt += 1

    return ImportErgebnis(angelegt=angelegt, uebersprungen=uebersprungen, fehler=fehler)
